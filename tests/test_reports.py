"""Testes da exportação para Excel (transações do ano e carteira para o IR)."""

from datetime import date, datetime

import pytest
from openpyxl import load_workbook

from stonks_ai import reports
from stonks_ai.database import db
from stonks_ai.models.investment import InvestmentOperation
from stonks_ai.models.transaction import Transaction
from stonks_ai.utils import portfolio


@pytest.fixture(autouse=True)
def export_to_tmp(tmp_path, monkeypatch):
    monkeypatch.setattr(reports, "export_dir", lambda: tmp_path)
    return tmp_path


def add_tx(day, desc, amount, kind="expense", category="Alimentação"):
    with db.session() as s:
        s.add(Transaction(description=desc, amount=amount, type=kind, category_name=category, date=datetime(*day)))


def rows(ws):
    return [list(r) for r in ws.iter_rows(min_row=2, values_only=True)]


def test_transactions_workbook_stops_today():
    add_tx((2026, 1, 5), "Salário", 5000, "income", "Salário")
    add_tx((2026, 1, 10), "Mercado", 300)
    add_tx((2026, 2, 3), "Uber", 40, category="Transporte")
    add_tx((2026, 3, 1), "Aluguel agendado", 1500, category="Moradia")   # depois de "hoje"
    add_tx((2025, 12, 31), "Ano passado", 99)

    path = reports.save(reports.transactions_workbook(2026, today=date(2026, 2, 20)), "t.xlsx")
    wb = load_workbook(path)

    assert wb.sheetnames == ["Resumo por mês", "Transações", "Despesas por categoria"]
    tx = rows(wb["Transações"])
    assert [r[1] for r in tx] == ["Salário", "Mercado", "Uber"]
    assert [r[5] for r in tx] == [5000, -300, -40]          # despesas negativas
    summary = rows(wb["Resumo por mês"])
    assert [r[0] for r in summary] == ["Jan/2026", "Fev/2026", "Total"]   # só até o mês de hoje
    assert summary[0][1:3] == [5000, 300]
    assert summary[-1][1] == "=SUM(B2:B3)"
    cats = {r[0]: r[1:3] for r in rows(wb["Despesas por categoria"])}
    assert cats["Alimentação"] == [300, 0]
    assert cats["Transporte"] == [0, 40]


def test_save_never_overwrites(export_to_tmp):
    wb = reports.transactions_workbook(2026)
    first = reports.save(wb, "x.xlsx")
    second = reports.save(wb, "x.xlsx")
    assert (first.name, second.name) == ("x.xlsx", "x (2).xlsx")


def add_op(kind, day, ticker="PETR4", qty=0.0, price=0.0, amount=0.0, source="manual", notes=None):
    with db.session() as s:
        s.add(InvestmentOperation(ticker=ticker, kind=kind, date=datetime(*day), quantity=qty, price=price,
                                  amount=amount, source=source, notes=notes))


def test_ir_workbook():
    add_op("buy", (2024, 3, 1), qty=100, price=30, source="b3", notes="XP INVESTIMENTOS")
    add_op("buy", (2025, 2, 1), qty=100, price=40, source="b3", notes="XP INVESTIMENTOS")
    add_op("sell", (2025, 6, 2), qty=50, price=50)
    add_op("dividend", (2025, 8, 20), amount=120)
    add_op("jcp", (2025, 9, 20), amount=30)
    add_op("jcp", (2024, 9, 20), amount=999)               # outro ano
    add_op("buy", (2025, 4, 1), ticker="HGLG11", qty=10, price=150, notes="anotação pessoal")
    add_op("sell", (2025, 11, 3), ticker="HGLG11", qty=10, price=160)
    add_op("income", (2025, 5, 14), ticker="HGLG11", amount=11)
    add_op("buy", (2024, 5, 2), ticker="VALE3", qty=10, price=60, source="b3", notes="RICO INVESTIMENTOS")
    add_op("sell", (2025, 3, 3), ticker="VALE3", qty=10, price=55)

    wb = load_workbook(reports.save(reports.ir_workbook(2025), "ir.xlsx"))
    assert wb.sheetnames == ["Leia-me", "Bens e direitos", "Proventos 2025", "Vendas 2025", "Vendas por mês"]

    bens = {r[0]: r for r in rows(wb["Bens e direitos"]) if r[0] != "Total"}
    petr = bens["PETR4"]
    assert petr[1] == "XP INVESTIMENTOS"
    assert petr[2:6] == [100, 3000, 150, 5250]               # custo, não valor de mercado
    assert petr[6] == 35
    assert petr[7] == "150 cotas de PETR4, ao preço médio de R$ 35,00, custodiadas em XP INVESTIMENTOS."
    assert "HGLG11" not in bens                               # comprado e vendido no mesmo ano
    # Observação digitada à mão não vira nome de corretora
    assert reports._brokers(portfolio.list_operations()) == {"PETR4": "XP INVESTIMENTOS", "VALE3": "RICO INVESTIMENTOS"}
    vale = bens["VALE3"]
    assert vale[2:6] == [10, 600, 0, 0]
    assert vale[7] == "Posição em VALE3 vendida em 2025."

    prov = {r[0]: r[1:4] for r in rows(wb["Proventos 2025"])}
    assert prov["PETR4"] == [120, 30, 0]
    assert prov["HGLG11"] == [0, 0, 11]

    vendas = rows(wb["Vendas 2025"])[:-1]
    assert [(v[1], v[3], v[4], v[5]) for v in vendas] == [
        ("VALE3", 550, 600, -50), ("PETR4", 2500, 1750, 750), ("HGLG11", 1600, 1500, 100)]
    assert [r[0] for r in rows(wb["Vendas por mês"])] == ["Mar/2025", "Jun/2025", "Nov/2025", "Total"]


def test_ir_workbook_with_crypto():
    add_op("buy", (2025, 1, 10), ticker="BTC", qty=0.004, price=500_000)
    add_op("sell", (2025, 6, 2), ticker="BTC", qty=0.001, price=600_000)
    add_op("buy", (2025, 2, 1), ticker="SHIB", qty=10_000_000, price=0.00008)
    add_op("buy", (2025, 2, 1), qty=10, price=30)
    add_op("sell", (2025, 6, 3), qty=10, price=35)

    wb = load_workbook(reports.save(reports.ir_workbook(2025), "ir.xlsx"))
    bens = {r[0]: r for r in rows(wb["Bens e direitos"]) if r[0] != "Total"}
    assert bens["BTC"][7] == "0,003 BTC (Bitcoin), custo total de aquisição de R$ 1.500,00."
    assert bens["SHIB"][6] == pytest.approx(0.00008)          # preço médio sem arredondar para centavos
    junho = rows(wb["Vendas por mês"])[0]
    assert junho == ["Jun/2025", 350, 50, 600, 100]            # bolsa e cripto em colunas separadas
    assert "35 mil" in " ".join(str(r[0]) for r in wb["Leia-me"].iter_rows(values_only=True))


def test_build_positions_collects_sales():
    sales = []
    portfolio.add_operation("ITSA4", "buy", date(2025, 1, 2), quantity=10, price=10)
    portfolio.build_positions(portfolio.list_operations(), sales=sales)
    assert sales == []
    portfolio.add_operation("ITSA4", "sell", date(2025, 2, 2), quantity=4, price=12)
    portfolio.build_positions(portfolio.list_operations(), sales=sales)
    assert sales == [{"date": date(2025, 2, 2), "ticker": "ITSA4", "quantity": 4, "known_quantity": 4,
                      "proceeds": 48, "cost": 40, "result": 8}]


def test_export_routes(export_to_tmp, monkeypatch):
    pytest.importorskip("fastapi")
    from fastapi.testclient import TestClient

    import desktop.api as api

    with TestClient(api.app, base_url="http://127.0.0.1") as c:
        r = c.post("/api/export/transactions", json={"year": 2026})
        assert r.status_code == 200
        assert r.json()["name"] == "stonks-transacoes-2026.xlsx"
        assert (export_to_tmp / "stonks-transacoes-2026.xlsx").is_file()
        assert c.post("/api/export/ir", json={"year": 2025}).json()["name"] == "stonks-carteira-ir-2025.xlsx"
        assert c.post("/api/export/ir", json={"year": 2999}).status_code == 422
        # A rota de mostrar na pasta não abre caminhos fora da pasta de exportação
        assert c.post("/api/export/show", json={"name": "..\\..\\Windows\\notepad.exe"}).status_code == 404
        assert c.post("/api/export/show", json={"name": "nao-existe.xlsx"}).status_code == 404
