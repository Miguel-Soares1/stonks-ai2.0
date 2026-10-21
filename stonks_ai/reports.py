"""
Exportação para Excel: transações do ano e relatório da carteira para o Imposto de Renda.

As planilhas são gravadas na pasta Downloads do usuário (ou em <dados>/exportacoes, se ela não
existir): a janela do app não baixa arquivos como um navegador.

O relatório do IR organiza o que a declaração pede (bens e direitos pelo custo, proventos e
vendas com o resultado pelo preço médio), mas não calcula imposto nem gera DARF.
"""

from collections import defaultdict
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Sequence

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from stonks_ai.collectors.crypto import ASSETS as CRYPTO_ASSETS, STABLECOINS, is_crypto
from stonks_ai.config import DATA_DIR
from stonks_ai.database import db
from stonks_ai.models.account import Account
from stonks_ai.models.investment import INCOME_KINDS
from stonks_ai.models.transaction import Transaction
from stonks_ai.utils.portfolio import EPS, build_positions, list_operations

MONTHS = ["Jan", "Fev", "Mar", "Abr", "Mai", "Jun", "Jul", "Ago", "Set", "Out", "Nov", "Dez"]
MONEY = '"R$" #,##0.00;[Red]-"R$" #,##0.00'
QTY = "#,##0.########"
DATE = "dd/mm/yyyy"
_HEADER_FILL = PatternFill("solid", fgColor="1F4FD8")


def export_dir() -> Path:
    downloads = Path.home() / "Downloads"
    folder = downloads if downloads.is_dir() else DATA_DIR / "exportacoes"
    folder.mkdir(parents=True, exist_ok=True)
    return folder


def save(wb: Workbook, filename: str) -> Path:
    """Grava sem sobrescrever: "nome.xlsx", depois "nome (2).xlsx" e assim por diante."""
    folder = export_dir()
    stem, suffix = Path(filename).stem, Path(filename).suffix
    path, n = folder / filename, 2
    while path.exists():
        path, n = folder / f"{stem} ({n}){suffix}", n + 1
    wb.save(path)
    return path


# ── Formatação ──────────────────────────────────────────────────────────

def _sheet(wb: Workbook, title: str, headers: Sequence[str], first: bool = False) -> Worksheet:
    ws = wb.active if first else wb.create_sheet()
    ws.title = title
    ws.append(list(headers))
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = _HEADER_FILL
        cell.alignment = Alignment(vertical="center")
    ws.freeze_panes = "A2"
    return ws


def _formats(ws: Worksheet, formats: Dict[int, str], widths: Dict[int, int]) -> None:
    """Formato por coluna (1 = A) para todas as linhas de dados, e larguras."""
    for row in ws.iter_rows(min_row=2):
        for col, fmt in formats.items():
            row[col - 1].number_format = fmt
    for col, width in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = width


def _total_row(ws: Worksheet, label: str, cols: Iterable[int], label_col: int = 1) -> None:
    """Linha de total com fórmulas SOMA, para a planilha continuar certa se o usuário editar."""
    last = ws.max_row
    row = [None] * ws.max_column
    row[label_col - 1] = label
    for col in cols:
        letter = get_column_letter(col)
        row[col - 1] = f"=SUM({letter}2:{letter}{last})"
    ws.append(row)
    for cell in ws[ws.max_row]:
        cell.font = Font(bold=True)


def _notes(wb: Workbook, title: str, lines: List[str]) -> None:
    ws = wb.create_sheet(title, 0)
    ws.column_dimensions["A"].width = 110
    for i, line in enumerate(lines, start=1):
        ws.cell(row=i, column=1, value=line).alignment = Alignment(wrap_text=True, vertical="top")
    ws["A1"].font = Font(bold=True, size=14)


# ── Transações do ano ───────────────────────────────────────────────────

def transactions_workbook(year: int, today: Optional[date] = None) -> Workbook:
    """Transações do ano até hoje, resumo por mês e despesas por categoria e mês."""
    today = today or date.today()
    start, end = datetime(year, 1, 1), datetime(year + 1, 1, 1)
    if year == today.year:  # o que está agendado para depois de hoje ainda não aconteceu
        end = datetime(today.year, today.month, today.day) + timedelta(days=1)
    with db.session() as s:
        accounts = {a.id: a.name for a in s.query(Account).all()}
        rows = [
            (t.date, t.description, t.category_name or "Sem categoria", accounts.get(t.account_id, ""),
             t.type, t.amount, t.notes or "")
            for t in s.query(Transaction)
            .filter(Transaction.date >= start, Transaction.date < end)
            .order_by(Transaction.date, Transaction.id).all()
        ]

    wb = Workbook()
    ws = _sheet(wb, "Resumo por mês", ["Mês", "Receitas", "Despesas", "Saldo"], first=True)
    income, expense = defaultdict(float), defaultdict(float)
    by_cat: Dict[str, Dict[int, float]] = defaultdict(lambda: defaultdict(float))
    for day, _, cat, _, kind, amount, _ in rows:
        if kind == "income":
            income[day.month] += amount
        else:
            expense[day.month] += amount
            by_cat[cat][day.month] += amount
    last_month = today.month if year == today.year else 12
    for m in range(1, last_month + 1):
        ws.append([f"{MONTHS[m - 1]}/{year}", income[m], expense[m], f"=B{ws.max_row + 1}-C{ws.max_row + 1}"])
    _total_row(ws, "Total", [2, 3, 4])
    _formats(ws, {2: MONEY, 3: MONEY, 4: MONEY}, {1: 14, 2: 16, 3: 16, 4: 16})

    ws = _sheet(wb, "Transações", ["Data", "Descrição", "Categoria", "Conta", "Tipo", "Valor", "Observações"])
    for day, desc, cat, account, kind, amount, notes in rows:
        signed = amount if kind == "income" else -amount
        ws.append([day.date(), desc, cat, account, "Receita" if kind == "income" else "Despesa", signed, notes])
    ws.auto_filter.ref = f"A1:G{max(ws.max_row, 1)}"
    _formats(ws, {1: DATE, 6: MONEY}, {1: 12, 2: 44, 3: 20, 4: 20, 5: 10, 6: 14, 7: 30})

    ws = _sheet(wb, "Despesas por categoria", ["Categoria", *MONTHS[:last_month], "Total"])
    total_col = last_month + 2
    for cat in sorted(by_cat, key=lambda c: -sum(by_cat[c].values())):
        r = ws.max_row + 1
        ws.append([cat, *[by_cat[cat][m] for m in range(1, last_month + 1)],
                   f"=SUM(B{r}:{get_column_letter(total_col - 1)}{r})"])
    _total_row(ws, "Total", range(2, total_col + 1))
    _formats(ws, {c: MONEY for c in range(2, total_col + 1)},
             {1: 24, **{c: 13 for c in range(2, total_col + 1)}})
    return wb


# ── Imposto de Renda ────────────────────────────────────────────────────

def _brl(value: float) -> str:
    """1234.5 -> "R$ 1.234,50" (para textos; as células usam o formato numérico)."""
    return "R$ " + f"{value:,.2f}".replace(",", "_").replace(".", ",").replace("_", ".")


def _crypto_qty(value: float) -> str:
    """0.0035 -> "0,0035" (até 8 casas, sem zeros à direita)."""
    return f"{value:.8f}".rstrip("0").rstrip(".").replace(".", ",")


def _brokers(ops) -> Dict[str, str]:
    """Corretoras por ativo, como vieram no extrato da B3 (observações manuais ficam de fora)."""
    found: Dict[str, set] = defaultdict(set)
    for op in ops:
        if op.source == "b3" and op.notes:
            found[op.ticker].add(op.notes)
    return {t: ", ".join(sorted(b)) for t, b in found.items()}


def ir_workbook(year: int) -> Workbook:
    """Posição em 31/12 (ano anterior e ano), proventos e vendas do ano, pelo custo de aquisição."""
    ops = list_operations()
    end_prev, end_year = date(year - 1, 12, 31), date(year, 12, 31)
    before = build_positions([op for op in ops if op.date.date() <= end_prev])
    sales: List[Dict[str, Any]] = []
    after = build_positions([op for op in ops if op.date.date() <= end_year], sales=sales)
    sales = [s for s in sales if s["date"].year == year]
    brokers = _brokers(ops)

    wb = Workbook()
    ws = _sheet(wb, "Bens e direitos", [
        "Ativo", "Corretora", f"Cotas em 31/12/{year - 1}", f"Situação em 31/12/{year - 1}",
        f"Cotas em 31/12/{year}", f"Situação em 31/12/{year}", "Preço médio", "Discriminação (sugestão)",
        "Atenção",
    ], first=True)
    tickers = sorted(t for t in set(before) | set(after)
                     if (before.get(t) and before[t].quantity > EPS) or (after.get(t) and after[t].quantity > EPS))
    for t in tickers:
        b, a = before.get(t), after.get(t)
        qb, cb = (b.quantity, round(b.cost, 2)) if b else (0.0, 0.0)
        qa, ca = (a.quantity, round(a.cost, 2)) if a else (0.0, 0.0)
        broker = brokers.get(t, "")
        if qa > EPS and is_crypto(t):
            # Preço médio de cripto pode ser R$ 0,00003 (SHIB): o custo total diz mais
            text = f"{_crypto_qty(qa)} {t} ({CRYPTO_ASSETS[t]}), custo total de aquisição de {_brl(ca)}."
        elif qa > EPS:
            text = f"{qa:g} cotas de {t}, ao preço médio de {_brl(a.avg_price)}"
            text += f", custodiadas em {broker}." if broker else "."
        else:
            text = f"Posição em {t} vendida em {year}."
        warn = "Faltam compras antigas: o custo pode estar errado." if (a and a.incomplete) or (b and b.incomplete) else ""
        ws.append([t, broker, qb, cb, qa, ca, round(a.avg_price, 8 if is_crypto(t) else 4) if a and qa > EPS else None, text, warn])
    _total_row(ws, "Total", [4, 6])
    _formats(ws, {3: QTY, 4: MONEY, 5: QTY, 6: MONEY, 7: MONEY},
             {1: 10, 2: 26, 3: 16, 4: 20, 5: 16, 6: 20, 7: 14, 8: 70, 9: 34})

    ws = _sheet(wb, f"Proventos {year}", ["Ativo", "Dividendos", "Juros sobre capital próprio", "Rendimentos (FII)", "Total"])
    income: Dict[str, Dict[str, float]] = defaultdict(lambda: defaultdict(float))
    for op in ops:
        if op.kind in INCOME_KINDS and op.date.year == year:
            income[op.ticker][op.kind] += op.amount
    for t in sorted(income):
        r = ws.max_row + 1
        ws.append([t, income[t]["dividend"], income[t]["jcp"], income[t]["income"], f"=SUM(B{r}:D{r})"])
    _total_row(ws, "Total", [2, 3, 4, 5])
    _formats(ws, {c: MONEY for c in range(2, 6)}, {1: 10, 2: 16, 3: 28, 4: 20, 5: 16})

    ws = _sheet(wb, f"Vendas {year}", ["Data", "Ativo", "Cotas", "Valor da venda", "Custo (preço médio)", "Resultado"])
    for s in sales:
        ws.append([s["date"], s["ticker"], s["quantity"], s["proceeds"], s["cost"], s["result"]])
    _total_row(ws, "Total", [4, 5, 6])
    _formats(ws, {1: DATE, 3: QTY, 4: MONEY, 5: MONEY, 6: MONEY}, {1: 12, 2: 10, 3: 10, 4: 16, 5: 20, 6: 16})

    # Bolsa e cripto à parte: cada uma tem a sua isenção mensal, sobre o total vendido no mês
    ws = _sheet(wb, "Vendas por mês", ["Mês", "Bolsa: total vendido", "Bolsa: resultado",
                                       "Cripto: total vendido", "Cripto: resultado"])
    by_month: Dict[int, List[float]] = defaultdict(lambda: [0.0, 0.0, 0.0, 0.0])
    for s in sales:
        col = 2 if is_crypto(s["ticker"]) else 0
        by_month[s["date"].month][col] += s["proceeds"]
        by_month[s["date"].month][col + 1] += s["result"]
    for m in sorted(by_month):
        ws.append([f"{MONTHS[m - 1]}/{year}", *by_month[m]])
    _total_row(ws, "Total", [2, 3, 4, 5])
    _formats(ws, {c: MONEY for c in range(2, 6)}, {1: 12, 2: 20, 3: 18, 4: 20, 5: 18})

    _notes(wb, "Leia-me", [
        f"Carteira para o Imposto de Renda {year + 1} (ano-calendário {year})",
        "Gerado pelo Stonks AI a partir das operações da Carteira. Confira com os informes de rendimentos "
        "das corretoras e da B3 antes de declarar; este relatório não substitui um contador.",
        "Bens e direitos: a situação em 31/12 é o CUSTO de aquisição (cotas × preço médio, com as taxas "
        "lançadas), não o valor de mercado. A discriminação é uma sugestão de texto; o CNPJ do ativo vem "
        "do informe da corretora.",
        "Proventos: dividendos e rendimentos de FII costumam ir em Rendimentos Isentos e Não Tributáveis; "
        "juros sobre capital próprio, em Rendimentos Sujeitos à Tributação Exclusiva. Os valores são os "
        "lançados (o extrato da B3 informa o líquido).",
        "Vendas: resultado pelo preço médio. Vendas de ações até R$ 20 mil no mês são isentas; FIIs, ETFs "
        "e BDRs não têm essa isenção. O app não separa esses tipos nem calcula o imposto (DARF).",
        "Criptomoedas: em Bens e direitos, grupo 08 (Criptoativos): código 01 para Bitcoin, 02 para as "
        f"outras moedas e 03 para stablecoins ({', '.join(sorted(STABLECOINS))}). Pela regra atual, vendas de "
        "cripto até R$ 35 mil no mês, somando todas as moedas, são isentas; a aba Vendas por mês separa bolsa "
        "e cripto. Confira se a regra mudou antes de declarar.",
        "Explicação das contas: docs/CARTEIRA.md no repositório do Stonks AI.",
    ])
    return wb
